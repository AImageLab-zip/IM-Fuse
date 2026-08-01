#!/usr/bin/env python3
"""Build the same reproduction-comparison workbook as
build_reproduction_comparison.py, but with one tab per mimosa_[size] model
only (mimosa_base, mimosa_large, mimosa_tiny, ...), instead of one tab per
every available model.

Reuses all discovery/parsing/writing logic from build_reproduction_comparison.py
and just restricts which models get a tab.
"""

from __future__ import annotations

import os
import statistics
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import build_reproduction_comparison as brc  # noqa: E402

import openpyxl
from openpyxl.utils import get_column_letter
from rich.progress import track

OUTPUT_FILE = "reproduction-comparison-mimosa-size.xlsx"

MIMOSA_SIZE_MODELS = {
    brc.normalize(n)
    for n in (
        "mimosa_base",
        "mimosa_gargantuan",
        "mimosa_huge",
        "mimosa_large",
        "mimosa_medium",
        "mimosa_micro",
        "mimosa_small",
        "mimosa_tiny",
    )
}


def main():
    legacy_path = os.path.join(brc.RESULTS_DIR, brc.LEGACY_FILE)
    legacy, display_names = brc.parse_legacy(legacy_path)
    models, internal_models = brc.discover(brc.RESULTS_DIR)

    models = {k: v for k, v in models.items() if k in MIMOSA_SIZE_MODELS}
    internal_models = {k: v for k, v in internal_models.items() if k in MIMOSA_SIZE_MODELS}

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    # completion_report() reports against every model with a config template;
    # restrict "expected" to the mimosa_[size] family so completion% and the
    # missing-fold list only reflect that family.
    report = {}
    for dsnum in brc.DATASETS:
        expected = [m for m in brc.expected_models(dsnum) if m in MIMOSA_SIZE_MODELS]
        found = 0
        missing = []
        for model_norm in expected:
            folds = models.get(model_norm, {}).get(brc.DATASETS[dsnum], {})
            have = sorted(set(folds.keys()) & set(brc.EXPECTED_FOLDS))
            found += len(have)
            for fold in brc.EXPECTED_FOLDS:
                if fold not in have:
                    missing.append((model_norm, fold))
        report[dsnum] = {
            "expected": expected,
            "found": found,
            "total": len(expected) * len(brc.EXPECTED_FOLDS),
            "missing": missing,
        }

    brc.write_summary_sheet(wb, report, display_names, models)

    # One tab per available mimosa_[size] model, sorted by display name.
    available = sorted(models.keys(), key=lambda n: display_names.get(n, n).lower())
    if not available:
        print("No available mimosa_[size] models found in", brc.RESULTS_DIR)
        return

    for model_norm in track(available, description="Writing model tabs"):
        display = display_names.get(model_norm, model_norm)
        ws = wb.create_sheet(title=brc.safe_sheet_title(display))
        ws.sheet_view.showGridLines = False
        folds_by_ds = models[model_norm]
        internal_folds_by_ds = internal_models.get(model_norm, {})

        next_row = 1
        stats_by_dataset = {}
        for dsnum, dataset in brc.DATASETS.items():
            next_row, region_stats = brc.write_table(
                ws, next_row, dataset, legacy, folds_by_ds, internal_folds_by_ds, display
            )
            stats_by_dataset[dsnum] = region_stats

        dice_scores, hd95_scores, mean_stds = {}, {}, {}
        for dsnum, stats in stats_by_dataset.items():
            dice_means = [v.get("Mean") for v in stats.get("Dice", {}).values() if v.get("Mean") is not None]
            hd95_means = [v.get("Mean") for v in stats.get("HD95", {}).values() if v.get("Mean") is not None]
            all_stds = [
                v.get("Std")
                for metric_stats in stats.values()
                for v in metric_stats.values()
                if v.get("Std") is not None
            ]
            dice_scores[dsnum] = statistics.fmean(dice_means) if dice_means else None
            hd95_scores[dsnum] = statistics.fmean(hd95_means) if hd95_means else None
            mean_stds[dsnum] = statistics.fmean(all_stds) if all_stds else None

        def _overall(per_dataset):
            present = [v for v in per_dataset.values() if v is not None]
            return statistics.fmean(present) if present else None

        dice_score_all = _overall(dice_scores)
        hd95_score_all = _overall(hd95_scores)
        mean_std_all = _overall(mean_stds)

        next_row += 1

        def _write_score_row(row, label, value):
            label_cell = ws.cell(row=row, column=1, value=label)
            label_cell.font = brc.BOLD
            value_cell = ws.cell(row=row, column=2, value=round(value, 4) if value is not None else None)
            value_cell.font, value_cell.alignment, value_cell.border = brc.BOLD, brc.CENTER, brc.BORDER
            return row + 1

        for dsnum, dataset in brc.DATASETS.items():
            next_row = _write_score_row(next_row, f"Dice score {dsnum}", dice_scores.get(dsnum))
            next_row = _write_score_row(next_row, f"HD95 score {dsnum}", hd95_scores.get(dsnum))
            next_row = _write_score_row(next_row, f"Mean Std {dsnum}", mean_stds.get(dsnum))
        next_row = _write_score_row(next_row, "Dice score all", dice_score_all)
        next_row = _write_score_row(next_row, "HD95 score all", hd95_score_all)
        next_row = _write_score_row(next_row, "Mean Std all", mean_std_all)

        ws.column_dimensions["A"].width = 5
        for i in range(2, 5):
            ws.column_dimensions[get_column_letter(i)].width = 5
        for i in range(5, ws.max_column + 1):
            letter = get_column_letter(i)
            if ws.column_dimensions[letter].width != 14:
                ws.column_dimensions[letter].width = 9
        ws.freeze_panes = "E4"

    out_path = os.path.join(brc.RESULTS_DIR, OUTPUT_FILE)
    wb.save(out_path)
    print(f"Wrote {out_path}")

    os.makedirs(brc.REPO_OUTPUTS_DIR, exist_ok=True)
    repo_out_path = os.path.join(brc.REPO_OUTPUTS_DIR, OUTPUT_FILE)
    wb.save(repo_out_path)
    print(f"Wrote {repo_out_path}")
    print(f"Tabs ({len(available)}): {', '.join(display_names.get(m, m) for m in available)}")
    for model_norm in available:
        for ds, folds in models[model_norm].items():
            print(f"  {display_names.get(model_norm, model_norm)} / {ds}: folds = {', '.join(sorted(folds))}")

    print()
    for dsnum, dataset in brc.DATASETS.items():
        info = report[dsnum]
        pct = 100 * info["found"] / info["total"] if info["total"] else 0.0
        print(f"{dataset} completion: {info['found']}/{info['total']} ({pct:.1f}%)")
        for model_norm, fold in info["missing"]:
            print(f"  missing: {display_names.get(model_norm, model_norm)} / {fold}")


if __name__ == "__main__":
    main()
